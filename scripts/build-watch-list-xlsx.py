#!/usr/bin/env python3
"""
Write alpoe-watch-list.xlsx: every watch page on the site, in the shape of the
Google Sheet, with the address of each page beside it.

    python3 scripts/build-watch-list-xlsx.py

The four columns are the sheet's own — Brand, Sub-Collection, Variant / Name,
Reference No. — so the file drops into Sheets looking like the one it came from.
The fifth is the page, as a real hyperlink.

── Where the rows come from ──

Two places, because the site has two:

  data/catalogue.csv + data/catalogue-extra.csv   the sheet's rows
  data/products.csv (type=watch)                  the hand-written listings

A page exists for every row in both. The brand pages hide a hand-written listing
when the sheet carries the same reference — a generic "Datejust 41 126334" tile
beside twenty-one specific ones is the same watch twice — but the page is still
there and still linked from search, so it is still a row here. The Source column
says which list each came from.

── Why the slugs are rebuilt here ──

lib/catalogue.ts derives a page's slug from its row; this script reproduces that
derivation rather than reading it back out of the build. That is a duplication,
and duplications drift — so the script ends by checking every URL it produced
against the URLs in the built sitemap, and refuses to write the file if the two
sets differ. Run `pnpm build` first; without a sitemap to check against the
script will say so and stop.
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SITE = "https://alpoelondon.com"
OUT = ROOT / "alpoe-watch-list.xlsx"
SITEMAP = ROOT / ".next" / "server" / "app" / "sitemap.xml.body"

# Mirrors WATCH_BRANDS in lib/taxonomy.ts. A brand absent here is a brand the
# site does not carry, and its rows are skipped exactly as the site skips them.
BRAND_SLUGS = {
    "rolex": "rolex",
    "patek philippe": "patek-philippe",
    "audemars piguet": "audemars-piguet",
    "richard mille": "richard-mille",
    "cartier": "cartier",
    "hublot": "hublot",
    "omega": "omega",
    "breitling": "breitling",
}


def slugify(s: str) -> str:
    """lib/catalogue.ts slugify(), character for character."""
    s = s.lower().replace("&", "and")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"^-|-$", "", s)


def reference_key(s: str) -> str:
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"^-|-$", "", s)


def read_rows(path: Path) -> list[list[str]]:
    with open(path, newline="", encoding="utf8") as f:
        return [row for row in csv.reader(f)]


def row_key(cols: list[str]) -> str:
    """lib/catalogue.ts rowKey(): brand / reference / variant."""
    pick = [cols[0] if len(cols) > 0 else "", cols[3] if len(cols) > 3 else "",
            cols[2] if len(cols) > 2 else ""]
    return "/".join(reference_key(c or "") for c in pick)


def catalogue_rows() -> list[dict]:
    """The sheet's rows, plus the extras it does not carry yet."""
    rows = read_rows(ROOT / "data" / "catalogue.csv")
    extra = read_rows(ROOT / "data" / "catalogue-extra.csv")[1:]  # header
    seen_rows = {row_key(r) for r in rows}
    rows = rows + [r for r in extra if len(r) >= 4 and row_key(r) not in seen_rows]

    out: list[dict] = []
    seen_slugs: set[str] = set()
    for cols in rows:
        brand = (cols[0] if len(cols) > 0 else "").strip()
        model = (cols[1] if len(cols) > 1 else "").strip()
        variant = (cols[2] if len(cols) > 2 else "").strip()
        reference = (cols[3] if len(cols) > 3 else "").strip()

        brand_slug = BRAND_SLUGS.get(brand.lower())
        # Header row, brand dividers, summary lines and blanks — skipped by the
        # site for the same reason: they are not listings.
        if not brand_slug:
            continue
        if not model and not variant and not reference:
            continue

        base = slugify(f"{model}-{variant}-{reference}")
        slug, n = base, 2
        while slug in seen_slugs:
            slug = f"{base}-{n}"
            n += 1
        seen_slugs.add(slug)

        out.append(
            {
                "brand": brand,
                "model": model,
                "variant": variant,
                "reference": reference,
                "url": f"{SITE}/watches/{brand_slug}/{slug}",
                "source": "Sheet",
            }
        )
    return out


def product_rows() -> list[dict]:
    """The hand-written listings in data/products.csv."""
    out: list[dict] = []
    with open(ROOT / "data" / "products.csv", newline="", encoding="utf8") as f:
        for r in csv.DictReader(f):
            if r["type"] != "watch":
                continue
            brand_slug = BRAND_SLUGS.get((r["brand"] or "").strip().lower())
            if not brand_slug:
                continue
            out.append(
                {
                    "brand": r["brand"],
                    "model": r["model"],
                    # The sheet's third column is the name a customer reads, and
                    # for these rows that is the title.
                    "variant": r["title"],
                    "reference": r["reference_number"],
                    "url": f"{SITE}/watches/{brand_slug}/{r['slug']}",
                    "source": "Listing",
                }
            )
    return out


def sitemap_watch_urls() -> set[str]:
    if not SITEMAP.exists():
        sys.exit(
            f"no built sitemap at {SITEMAP.relative_to(ROOT)} — run `pnpm build` first, "
            "so the URLs in this file can be checked against the ones the site serves"
        )
    body = SITEMAP.read_text(encoding="utf8")
    locs = re.findall(r"<loc>([^<]+)</loc>", body)
    return {u for u in locs if "/watches/" in u and u.count("/") >= 5}


def main() -> None:
    rows = catalogue_rows() + product_rows()

    # The duplication guard. If the slug derivation above has drifted from
    # lib/catalogue.ts, these two sets stop matching and the file is not written.
    ours = {r["url"] for r in rows}
    theirs = sitemap_watch_urls()
    if ours != theirs:
        missing = sorted(theirs - ours)[:5]
        extra = sorted(ours - theirs)[:5]
        sys.exit(
            "URL sets differ — the slug derivation here no longer matches "
            f"lib/catalogue.ts.\n  in sitemap, not here ({len(theirs - ours)}): {missing}"
            f"\n  here, not in sitemap ({len(ours - theirs)}): {extra}"
        )

    # Brands in the order the site lists them, then rows as their source has them.
    order = {slug: i for i, slug in enumerate(BRAND_SLUGS.values())}
    rows.sort(key=lambda r: order.get(BRAND_SLUGS.get(r["brand"].lower(), ""), 99))

    wb = Workbook()
    ws = wb.active
    ws.title = "Watches"

    headers = ["Brand", "Sub-Collection", "Variant / Name", "Reference No.", "Page", "Source"]
    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="171312")
    for i in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")

    for r in rows:
        ws.append([r["brand"], r["model"], r["variant"], r["reference"], r["url"], r["source"]])
        cell = ws.cell(row=ws.max_row, column=5)
        # A real hyperlink rather than a HYPERLINK() formula, so the address is
        # still readable and copyable in the cell after Sheets imports it.
        cell.hyperlink = r["url"]
        cell.font = Font(color="0563C1", underline="single")

    widths = [18, 24, 46, 22, 62, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    wb.save(OUT)
    by_source: dict[str, int] = {}
    for r in rows:
        by_source[r["source"]] = by_source.get(r["source"], 0) + 1
    print(
        f"{OUT.name} written: {len(rows)} watches "
        f"({', '.join(f'{v} from {k.lower()}' for k, v in sorted(by_source.items()))}), "
        "every URL checked against the built sitemap"
    )


if __name__ == "__main__":
    main()
